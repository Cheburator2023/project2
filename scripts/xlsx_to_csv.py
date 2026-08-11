#!/usr/bin/env python3
"""Convert a multi-sheet XLSX workbook into one CSV file per sheet."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd


INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def sanitize_sheet_name(name: str) -> str:
    cleaned = INVALID_FILENAME_CHARS.sub("_", name.strip())
    return cleaned or "sheet"


def unique_filename(base: str, used: set[str]) -> str:
    candidate = f"{base}.csv"
    if candidate not in used:
        used.add(candidate)
        return candidate

    index = 2
    while True:
        candidate = f"{base}_{index}.csv"
        if candidate not in used:
            used.add(candidate)
            return candidate
        index += 1


def _sheet_visibility(xlsx_path: Path) -> dict[str, str]:
    """Return sheet_name -> state (visible|hidden|veryHidden)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    wb = load_workbook(xlsx_path, read_only=True, data_only=False)
    visibility: dict[str, str] = {}
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            state = getattr(ws, "sheet_state", None) or "visible"
            visibility[name] = str(state)
    finally:
        wb.close()
    return visibility


def convert_xlsx_to_csv(xlsx_path: Path, output_dir: Path) -> list[Path]:
    if not xlsx_path.is_file():
        raise FileNotFoundError(f"XLSX not found: {xlsx_path}")

    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = pd.ExcelFile(xlsx_path, engine="openpyxl")
    visibility = _sheet_visibility(xlsx_path)
    written: list[Path] = []
    used_names: set[str] = set()
    skipped_hidden: list[str] = []

    for sheet_name in workbook.sheet_names:
        state = visibility.get(sheet_name, "visible").lower()
        if state in {"hidden", "veryhidden"}:
            skipped_hidden.append(f"{sheet_name!r} ({state})")
            print(f"Skip hidden sheet {sheet_name!r} (state={state})")
            continue

        df = pd.read_excel(workbook, sheet_name=sheet_name, header=None, dtype=object)
        filename = unique_filename(sanitize_sheet_name(sheet_name), used_names)
        csv_path = output_dir / filename
        df.to_csv(csv_path, index=False, header=False, encoding="utf-8-sig")
        written.append(csv_path)
        print(
            f"Wrote {csv_path.name} ({len(df)} rows × {df.shape[1]} cols) "
            f"from sheet {sheet_name!r}"
        )

    if skipped_hidden:
        print(f"Skipped {len(skipped_hidden)} hidden sheet(s): {', '.join(skipped_hidden)}")
    elif not visibility:
        print(
            "Note: sheet visibility unavailable (openpyxl missing); "
            "exported all sheets from pandas."
        )

    return written


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    default_xlsx = Path(__file__).resolve().parent / (
        "Мастер версия_05.08.26"
        "/Файл для для загрузки 87 - Мастер версия_05.08.26.xlsx"
    )

    parser = argparse.ArgumentParser(
        description="Convert each sheet of an XLSX workbook into a separate CSV file."
    )
    parser.add_argument(
        "xlsx",
        nargs="?",
        type=Path,
        default=default_xlsx,
        help=f"Path to the XLSX file (default: {default_xlsx.name})",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        type=Path,
        default=None,
        help="Directory for CSV files (default: <xlsx_stem>_csv next to the XLSX)",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    xlsx_path = args.xlsx.resolve()
    output_dir = (
        args.output_dir.resolve()
        if args.output_dir is not None
        else xlsx_path.parent / f"{xlsx_path.stem}_csv"
    )

    try:
        written = convert_xlsx_to_csv(xlsx_path, output_dir)
    except Exception as exc:  # noqa: BLE001 - CLI entrypoint
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print(f"Done: {len(written)} CSV file(s) in {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
